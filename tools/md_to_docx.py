"""Converts one of the project's Markdown documents to a Word file (python-docx).

Handles what the project documents use: headings, paragraphs, bullet and numbered lists,
pipe tables, block quotes, images (local paths) and inline **bold**, *italic*, `code` and links.
Usage:  python tools/md_to_docx.py SwarmRide_Bharat_Pitch.md SwarmRide_Bharat_Pitch.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x13, 0x23, 0x3A)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("[") and "](" in part:
            label = part[1:part.index("](")]
            run = par.add_run(label)
            run.font.color.rgb = RGBColor(0x1F, 0x6F, 0x66)
            run.underline = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_path, docx_path):
    md_path = Path(md_path)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or s.startswith("<") or s == "---":
            i += 1
            continue
        if s.startswith("```"):                                    # code block
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(block))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            h = doc.add_heading(level=len(m.group(1)) - 1)              # # title, ## heading 1, ...
            add_inline(h, m.group(2))
            for r in h.runs:
                r.font.color.rgb = NAVY
            i += 1
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:                                                      # image
            img = (md_path.parent / m.group(2)).resolve()
            if img.exists() and img.suffix.lower() in (".png", ".jpg", ".jpeg"):
                doc.add_picture(str(img), width=Inches(6.3))
                cap = doc.add_paragraph()
                cap.add_run(m.group(1)).italic = True
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[i + 1].strip()):
            head = cells(s)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(cells(lines[i]))
                i += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(head))
            t.style = "Light Grid Accent 1"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, txt in enumerate(head):
                add_inline(t.rows[0].cells[c].paragraphs[0], txt)
                for r in t.rows[0].cells[c].paragraphs[0].runs:
                    r.bold = True
            for r_i, row in enumerate(rows, start=1):
                for c in range(len(head)):
                    add_inline(t.rows[r_i].cells[c].paragraphs[0], row[c] if c < len(row) else "")
            doc.add_paragraph()
            continue
        if s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, s.lstrip("> ").strip())
            i += 1
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if m:
            p = doc.add_paragraph(style="List Number" if m.group(2)[0].isdigit() else "List Bullet")
            add_inline(p, m.group(3))
            i += 1
            continue
        p = doc.add_paragraph()                                    # paragraph (joins wrapped lines)
        text = s
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||!\[|```|>|[-*] |\d+\. )", lines[i].strip()):
            text += " " + lines[i].strip()
            i += 1
        add_inline(p, text)
    doc.save(docx_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: python tools/md_to_docx.py input.md output.docx")
    convert(sys.argv[1], sys.argv[2])
    print("wrote", sys.argv[2])
